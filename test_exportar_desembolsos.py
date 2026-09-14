import csv
import importlib.util
import json
from pathlib import Path

_SPEC = importlib.util.spec_from_file_location(
    "exportar_desembolsos",
    Path(__file__).resolve().parent / "exportar-desembolsos.py",
)
des = importlib.util.module_from_spec(_SPEC)
_SPEC.loader.exec_module(des)


CAMPOS = [
    "Crear hora", "Casa de apuestas", "ID Padre", "ID de usuario", "Usuario",
    "Tipo", "ID de transacción", "Nombre de usuario del emisor", "Moneda",
    "Ingresos", "Estado", "Total", "Comisión", "Saldo", "Saldo actual",
    "Billeteras", "Tipo de transacción", "grupo causal", "causal",
    "producto causal", "Descripción", "Nota", "Dirección IP",
]


def _fila(**kw):
    fila = {c: "" for c in CAMPOS}
    fila.update(kw)
    return fila


def _escribir_csv(path, filas):
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=CAMPOS)
        writer.writeheader()
        writer.writerows(filas)


# --------------------------------------------------------------------------- #
# parse_money
# --------------------------------------------------------------------------- #
def test_parse_money_formatos():
    assert des.parse_money("-150.00") == -150.0
    assert des.parse_money("$1,234.56") == 1234.56
    assert des.parse_money("1.234,56") == 1234.56
    assert des.parse_money("75,50") == 75.5
    assert des.parse_money("") == 0.0
    assert des.parse_money(None) == 0.0
    assert des.parse_money("n/a") == 0.0


# --------------------------------------------------------------------------- #
# es_desembolso
# --------------------------------------------------------------------------- #
def _cfg():
    return set(des.PAYMENT_PRODUCTS), tuple(des.WITHDRAWAL_TERMS)


def test_retiro_de_caja_es_desembolso():
    productos, terminos = _cfg()
    row = _fila(**{"producto causal": "Payments", "Tipo de transacción": "Withdraw",
                   "grupo causal": "Retiro", "Total": "-150.00"})
    assert des.es_desembolso(row, "pagos", productos, terminos) is True


def test_apuesta_casino_no_es_desembolso():
    productos, terminos = _cfg()
    # "Withdraw" de casino: mismo tipo de transacción, pero producto Casino.
    row = _fila(**{"producto causal": "Casino", "Tipo de transacción": "Withdraw",
                   "grupo causal": "Apuesta", "causal": "Bet", "Total": "-0.10"})
    assert des.es_desembolso(row, "pagos", productos, terminos) is False


def test_ganancia_casino_no_es_desembolso():
    productos, terminos = _cfg()
    row = _fila(**{"producto causal": "Casino", "Tipo de transacción": "Deposit",
                   "grupo causal": "Ganancia", "Total": "0.20"})
    assert des.es_desembolso(row, "pagos", productos, terminos) is False


def test_modo_amplio_incluye_terminos_en_cualquier_producto():
    productos, terminos = _cfg()
    row = _fila(**{"producto causal": "Miscellaneous", "grupo causal": "Retiro",
                   "Descripción": "Desembolso manual", "Total": "-20.00"})
    # En modo pagos no cuenta (producto no es de caja); en amplio sí.
    assert des.es_desembolso(row, "pagos", productos, terminos) is False
    assert des.es_desembolso(row, "amplio", productos, terminos) is True


def test_modo_withdraw_es_crudo():
    productos, terminos = _cfg()
    row = _fila(**{"producto causal": "Casino", "Tipo de transacción": "Withdraw"})
    assert des.es_desembolso(row, "withdraw", productos, terminos) is True


# --------------------------------------------------------------------------- #
# cargar_usuarios + enriquecimiento
# --------------------------------------------------------------------------- #
def test_cargar_usuarios(tmp_path):
    users = tmp_path / "usuarios-historico.json"
    users.write_text(json.dumps({
        "usuarios": {
            "788": {
                "usuario": "50300000000", "tipo": "Player Online", "canal": "online",
                "estado_actual": "activo", "dias_activos": 6, "transacciones": 2583,
                "apuesta_total": 1380.25, "ganancia_neta": 5.39,
                "primer_dia": "2026-09-09", "ultimo_dia": "2026-09-14",
            }
        }
    }), encoding="utf-8")
    act = tmp_path / "actividad-usuarios.json"
    act.write_text(json.dumps({
        "top_apostadores": [
            {"id_usuario": "788", "productos": ["Casino", "Payments"],
             "ultima_actividad": "2026-09-14 12:22:25"}
        ]
    }), encoding="utf-8")

    perfiles = des.cargar_usuarios(users, act)
    assert perfiles["788"]["hist_canal"] == "online"
    assert perfiles["788"]["hist_apuesta_total"] == 1380.25
    assert perfiles["788"]["hist_productos"] == "Casino, Payments"
    assert perfiles["788"]["hist_ultima_actividad"] == "2026-09-14 12:22:25"


def test_iter_desembolsos_enriquece_y_filtra(tmp_path):
    csv_path = tmp_path / "tx.csv"
    _escribir_csv(csv_path, [
        _fila(**{"ID de usuario": "788", "Usuario": "50300000000",
                 "producto causal": "Payments", "Tipo de transacción": "Withdraw",
                 "grupo causal": "Retiro", "Estado": "Pendiente", "Total": "-150.00",
                 "Tipo": "Player Online"}),
        _fila(**{"ID de usuario": "788", "Usuario": "50300000000",
                 "producto causal": "Casino", "Tipo de transacción": "Withdraw",
                 "grupo causal": "Apuesta", "Total": "-0.10"}),  # apuesta, se descarta
        _fila(**{"ID de usuario": "999", "Usuario": "50399999999",
                 "producto causal": "Banco", "Tipo de transacción": "Withdraw",
                 "causal": "Cash out", "Estado": "Aprobado", "Total": "-75.50"}),
    ])
    perfiles = {"788": {"hist_canal": "online", "hist_estado": "activo",
                        "hist_apuesta_total": 1380.25}}
    productos, terminos = _cfg()
    registros = list(des.iter_desembolsos(csv_path, perfiles, "pagos", productos, terminos))

    assert len(registros) == 2  # se excluye la apuesta de casino
    r0 = registros[0]
    assert r0["Monto solicitado"] == 150.0
    assert r0["Total (bruto)"] == -150.0
    assert r0["Cliente · canal"] == "online"        # enriquecido desde el perfil
    assert r0["Cliente · apostado total"] == 1380.25
    # Usuario sin perfil: campos de cliente vacíos, sin romper.
    r1 = registros[1]
    assert r1["Monto solicitado"] == 75.5
    assert r1["Cliente · canal"] == ""


def test_iter_desembolsos_extra_filter(tmp_path):
    csv_path = tmp_path / "tx.csv"
    _escribir_csv(csv_path, [
        _fila(**{"ID de usuario": "1", "Usuario": "aaa", "producto causal": "Payments",
                 "grupo causal": "Retiro", "Total": "-10"}),
        _fila(**{"ID de usuario": "2", "Usuario": "bbb", "producto causal": "Payments",
                 "grupo causal": "Retiro", "Total": "-20"}),
    ])
    productos, terminos = _cfg()
    solo_bbb = list(des.iter_desembolsos(
        csv_path, {}, "pagos", productos, terminos,
        extra_filter=lambda row: row.get("Usuario") == "bbb",
    ))
    assert len(solo_bbb) == 1
    assert solo_bbb[0]["Usuario"] == "bbb"


# --------------------------------------------------------------------------- #
# Generación del Excel (extremo a extremo)
# --------------------------------------------------------------------------- #
def test_generar_excel(tmp_path, monkeypatch):
    csv_path = tmp_path / "tx.csv"
    _escribir_csv(csv_path, [
        _fila(**{"ID de usuario": "788", "Usuario": "50300000000",
                 "producto causal": "Payments", "grupo causal": "Retiro",
                 "Estado": "Pendiente", "Total": "-150.00", "Tipo": "Player Online"}),
        _fila(**{"ID de usuario": "999", "Usuario": "50399999999",
                 "producto causal": "Banco", "causal": "Cash out",
                 "Estado": "Aprobado", "Total": "-75.50", "Tipo": "Player Retail"}),
    ])
    # JSON de usuarios apuntando a tmp_path.
    users = tmp_path / "usuarios-historico.json"
    users.write_text(json.dumps({"usuarios": {"788": {"canal": "online",
                     "estado_actual": "activo", "apuesta_total": 1380.25}}}),
                     encoding="utf-8")
    monkeypatch.setattr(des, "USUARIOS_JSON", users)
    monkeypatch.setattr(des, "ACTIVIDAD_JSON", tmp_path / "no-existe.json")

    salida = tmp_path / "desembolsos.xlsx"
    productos, terminos = _cfg()
    ruta, n = des.generar(csv_path, salida, "pagos", productos, terminos)
    assert n == 2
    assert ruta.exists()

    from openpyxl import load_workbook
    wb = load_workbook(ruta)
    assert wb.sheetnames == ["Desembolsos", "Por usuario", "Resumen"]
    ws = wb["Desembolsos"]
    cols = [c.value for c in ws[1]]
    assert "Monto solicitado" in cols
    assert "Cliente · canal" in cols
    assert ws.max_row - 1 == 2  # dos filas de detalle
