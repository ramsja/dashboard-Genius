#!/usr/bin/env python3
"""Exporta a Excel las solicitudes de desembolso (retiros) enriquecidas con el JSON de usuarios.

Del módulo de transacciones (los CSV de ``descargas/``) filtra únicamente los
movimientos que son **retiros / desembolsos de caja** (dinero que el cliente
pide sacar de la plataforma) y construye un Excel con:

  * Hoja ``Desembolsos``  : una fila por solicitud, con cada dato de la
    transacción MÁS los datos del cliente mapeados desde
    ``dashboard/data/usuarios-historico.json`` (canal, tipo, estado,
    días activos, apostado histórico, ganancia neta, etc.).
  * Hoja ``Por usuario``  : resumen agregado por cliente.
  * Hoja ``Resumen``      : KPIs generales (por canal, estado y producto).

Ojo con la nomenclatura del Back Office: la columna ``Tipo de transacción``
usa "Withdraw"/"Deposit" también para el flujo de casino (apuesta/ganancia).
Por eso, por defecto, un retiro real se detecta por el *producto de caja*
(Payments/Banco/…) o por términos explícitos de retiro en grupo/causal, no por
el "Withdraw" del casino. El criterio es configurable (``--modo`` y variables
de entorno).

Uso
---
    python exportar-desembolsos.py                      # usa el CSV más reciente de descargas/
    python exportar-desembolsos.py ruta/al.csv          # un CSV concreto
    python exportar-desembolsos.py --salida retiros.xlsx
    python exportar-desembolsos.py --modo amplio        # incluye términos de retiro en cualquier producto
    python exportar-desembolsos.py --watch              # regenera en tiempo real cuando cambia el CSV
    python exportar-desembolsos.py --watch --intervalo 15

Variables de entorno (opcionales, amplían/afinan la detección):
    DESEMBOLSO_PRODUCTOS   productos de caja extra, separados por coma
    DESEMBOLSO_TERMINOS    términos de retiro extra, separados por coma
"""
from __future__ import annotations

import argparse
import csv
import os
import sys
import time
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Iterator

BASE_DIR = Path(__file__).resolve().parent
DEFAULT_CSV_DIR = BASE_DIR / "descargas"
USUARIOS_JSON = BASE_DIR / "dashboard" / "data" / "usuarios-historico.json"
ACTIVIDAD_JSON = BASE_DIR / "dashboard" / "data" / "actividad-usuarios.json"
DEFAULT_OUTPUT = BASE_DIR / "reportes" / "desembolsos.xlsx"

# Productos "de caja" (cashier) que representan dinero real entrando/saliendo,
# no jugadas de casino. Se comparan sin acentos y en minúsculas.
PAYMENT_PRODUCTS = {
    "payments", "payment", "pago", "pagos", "banco", "bank", "banca",
    "cajero", "cashier", "cash", "caja", "retiros", "withdrawals",
}

# Términos que denotan un retiro/desembolso explícito en grupo/causal/descripción.
WITHDRAWAL_TERMS = (
    "retiro", "retirar", "desembolso", "reintegro", "cobro", "cash out",
    "cashout", "payout", "withdraw", "withdrawal", "solicitud de retiro",
    "solicitud de pago", "pago a cliente", "cash-out",
)


# --------------------------------------------------------------------------- #
# Utilidades de normalización de columnas (accent-insensitive, como el resto).
# --------------------------------------------------------------------------- #
def _norm_key(text: str) -> str:
    nfkd = unicodedata.normalize("NFKD", text or "")
    return nfkd.encode("ascii", "ignore").decode("ascii").casefold().strip()


def _get(row: dict[str, Any], *names: str) -> str:
    """Obtiene un valor de la fila buscando por varios nombres, sin acentos."""
    lookup = {_norm_key(k): k for k in row.keys()}
    for name in names:
        key = lookup.get(_norm_key(name))
        if key is not None:
            value = row.get(key)
            return "" if value is None else str(value).strip()
    return ""


def parse_money(value: Any) -> float:
    if value is None:
        return 0.0
    text = str(value).replace("$", "").replace(" ", "").strip()
    if not text:
        return 0.0
    # Formatos "1,234.56" y "1.234,56": si hay ambos, la última es decimal.
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        # Coma sola: decimal europeo si va seguida de 1-2 dígitos al final.
        if len(text.split(",")[-1]) in (1, 2):
            text = text.replace(",", ".")
        else:
            text = text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return 0.0


# --------------------------------------------------------------------------- #
# Detección de desembolsos.
# --------------------------------------------------------------------------- #
def _contains(text: str, terms: Iterable[str]) -> bool:
    return any(term in text for term in terms)


def es_desembolso(row: dict[str, Any], modo: str,
                  productos: set[str], terminos: tuple[str, ...]) -> bool:
    """Decide si una fila del CSV es una solicitud de retiro/desembolso.

    modo:
      * "pagos"    (defecto): sólo movimientos de caja (producto de pago) que
                    salen hacia el cliente. Es el retiro "real".
      * "amplio":  lo anterior + cualquier fila con términos de retiro en
                    grupo/causal/descripción, sea cual sea el producto.
      * "withdraw": toda fila cuyo 'Tipo de transacción' sea Withdraw (crudo;
                    incluye apuestas de casino). Sólo para diagnóstico.
    """
    tipo_tx = _norm_key(_get(row, "Tipo de transacción", "Tipo de transaccion", "transaction_type"))
    grupo = _norm_key(_get(row, "grupo causal", "grupo_causal", "causal_group"))
    causal = _norm_key(_get(row, "causal"))
    producto = _norm_key(_get(row, "producto causal", "producto_causal", "causal_product"))
    descripcion = _norm_key(_get(row, "Descripción", "Descripcion", "description"))
    blob = " ".join((grupo, causal, descripcion))

    if modo == "withdraw":
        return "withdraw" in tipo_tx or "retiro" in tipo_tx

    tiene_termino = _contains(blob, terminos)
    es_producto_caja = producto in productos or _contains(producto, productos)

    if modo == "amplio":
        if tiene_termino:
            return True

    if not es_producto_caja:
        return False

    # Producto de caja: aceptar si es dirección de salida (retiro) o si el
    # grupo/causal lo dice, o si el monto es negativo (dinero que sale).
    if "withdraw" in tipo_tx or "retiro" in tipo_tx or tiene_termino:
        return True
    total = parse_money(_get(row, "Total", "total"))
    return total < 0


# --------------------------------------------------------------------------- #
# Carga y mapeo del JSON de usuarios.
# --------------------------------------------------------------------------- #
def cargar_usuarios(path: Path = USUARIOS_JSON,
                    actividad: Path = ACTIVIDAD_JSON) -> dict[str, dict[str, Any]]:
    """Devuelve {id_usuario: {campos del perfil}} desde el JSON del dashboard."""
    import json

    perfiles: dict[str, dict[str, Any]] = {}
    if path.exists():
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            data = {}
        for uid, info in (data.get("usuarios") or {}).items():
            if not isinstance(info, dict):
                continue
            perfiles[str(uid)] = {
                "hist_usuario": info.get("usuario", ""),
                "hist_tipo": info.get("tipo", ""),
                "hist_canal": info.get("canal", ""),
                "hist_estado": info.get("estado_actual", ""),
                "hist_primer_dia": info.get("primer_dia", ""),
                "hist_ultimo_dia": info.get("ultimo_dia", ""),
                "hist_dias_activos": info.get("dias_activos", ""),
                "hist_transacciones": info.get("transacciones", ""),
                "hist_apuesta_total": info.get("apuesta_total", ""),
                "hist_ganancia_neta": info.get("ganancia_neta", ""),
            }

    # Enriquecer con productos usados y última actividad (actividad-usuarios.json).
    if actividad.exists():
        try:
            act = json.loads(actividad.read_text(encoding="utf-8"))
        except Exception:
            act = {}
        for grupo in ("top_apostadores", "usuarios", "top"):
            filas = act.get(grupo)
            if not isinstance(filas, list):
                continue
            for item in filas:
                if not isinstance(item, dict):
                    continue
                uid = str(item.get("id_usuario") or "")
                if not uid:
                    continue
                perfil = perfiles.setdefault(uid, {})
                if item.get("productos"):
                    perfil["hist_productos"] = ", ".join(map(str, item["productos"]))
                if item.get("ultima_actividad"):
                    perfil["hist_ultima_actividad"] = item["ultima_actividad"]
    return perfiles


# Campos del cliente que se añaden a cada fila de desembolso (orden y etiqueta).
PERFIL_COLUMNS = [
    ("hist_tipo", "Cliente · tipo"),
    ("hist_canal", "Cliente · canal"),
    ("hist_estado", "Cliente · estado"),
    ("hist_dias_activos", "Cliente · días activos"),
    ("hist_transacciones", "Cliente · nº transacciones"),
    ("hist_apuesta_total", "Cliente · apostado total"),
    ("hist_ganancia_neta", "Cliente · ganancia neta"),
    ("hist_primer_dia", "Cliente · primer día"),
    ("hist_ultimo_dia", "Cliente · último día"),
    ("hist_ultima_actividad", "Cliente · última actividad"),
    ("hist_productos", "Cliente · productos"),
]

# Campos de la transacción que se extraen del CSV (etiqueta -> nombres CSV).
TX_COLUMNS: list[tuple[str, tuple[str, ...]]] = [
    ("Fecha/hora", ("Crear hora", "crear_hora", "created_at")),
    ("ID de usuario", ("ID de usuario", "id_usuario", "user_id")),
    ("Usuario", ("Usuario", "usuario", "username")),
    ("ID de transacción", ("ID de transacción", "ID de transaccion", "id_transaccion")),
    ("Tipo de transacción", ("Tipo de transacción", "Tipo de transaccion")),
    ("Grupo causal", ("grupo causal", "grupo_causal")),
    ("Causal", ("causal",)),
    ("Producto causal", ("producto causal", "producto_causal")),
    ("Descripción", ("Descripción", "Descripcion", "description")),
    ("Estado", ("Estado", "estado", "status")),
    ("Conexión (Tipo)", ("Tipo", "tipo")),
    ("Moneda", ("Moneda", "moneda")),
    ("Billeteras", ("Billeteras", "billeteras", "wallet")),
    ("Saldo", ("Saldo", "saldo")),
    ("Saldo actual", ("Saldo actual", "saldo_actual")),
    ("Comisión", ("Comisión", "Comision", "comision")),
    ("Ingresos", ("Ingresos", "ingresos")),
    ("Casa de apuestas", ("Casa de apuestas", "casa_apuestas", "site")),
    ("Dirección IP", ("Dirección IP", "Direccion IP", "direccion_ip")),
    ("Emisor", ("Nombre de usuario del emisor", "nombre_usuario_emisor")),
    ("Nota", ("Nota", "nota")),
]

MONEY_LABELS = {
    "Monto solicitado", "Total (bruto)", "Saldo", "Saldo actual", "Comisión",
    "Ingresos", "Cliente · apostado total", "Cliente · ganancia neta",
}


def _open_csv(path: Path):
    return path.open("r", encoding="utf-8-sig", newline="", errors="replace")


def iter_desembolsos(csv_path: Path, perfiles: dict[str, dict[str, Any]],
                     modo: str, productos: set[str],
                     terminos: tuple[str, ...],
                     extra_filter: Any = None) -> Iterator[dict[str, Any]]:
    """Genera un dict por cada fila de desembolso, ya enriquecida.

    ``extra_filter`` (opcional) es un callable ``(row) -> bool`` que permite
    aplicar filtros adicionales (fecha, búsqueda) además del de desembolso;
    lo usa el visor para respetar los filtros de la interfaz.
    """
    with _open_csv(csv_path) as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            if not es_desembolso(row, modo, productos, terminos):
                continue
            if extra_filter is not None and not extra_filter(row):
                continue
            registro: dict[str, Any] = {}
            for label, names in TX_COLUMNS:
                registro[label] = _get(row, *names)
            total = parse_money(_get(row, "Total", "total"))
            registro["Total (bruto)"] = total
            registro["Monto solicitado"] = abs(total)
            # Enriquecimiento por id de usuario.
            uid = _get(row, "ID de usuario", "id_usuario", "user_id")
            perfil = perfiles.get(uid, {})
            for key, label in PERFIL_COLUMNS:
                registro[label] = perfil.get(key, "")
            registro["_uid"] = uid
            yield registro


# Orden final de columnas en la hoja de detalle (transacción + montos + cliente).
def columnas_detalle() -> list[str]:
    cols = ["Fecha/hora", "ID de usuario", "Usuario", "ID de transacción",
            "Monto solicitado", "Total (bruto)", "Moneda", "Estado",
            "Tipo de transacción", "Grupo causal", "Causal", "Producto causal",
            "Descripción", "Conexión (Tipo)", "Billeteras", "Saldo",
            "Saldo actual", "Comisión", "Ingresos", "Casa de apuestas",
            "Dirección IP", "Emisor", "Nota"]
    cols += [label for _, label in PERFIL_COLUMNS]
    return cols


# --------------------------------------------------------------------------- #
# Escritura del Excel.
# --------------------------------------------------------------------------- #
def _to_number(value: Any) -> Any:
    """Convierte importes de texto a número para las celdas de dinero."""
    if isinstance(value, (int, float)):
        return value
    if value in (None, ""):
        return None
    return parse_money(value)


def construir_libro(registros: list[dict[str, Any]], origen: Path, modo: str):
    """Construye y devuelve el ``Workbook`` de openpyxl (sin guardarlo)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - depende del entorno
        raise SystemExit(
            "Falta 'openpyxl'. Instálalo con:  python -m pip install openpyxl"
        ) from exc

    book = Workbook()

    header_fill = PatternFill("solid", fgColor="0D5654")
    header_font = Font(color="FFFDF8", bold=True)
    title_font = Font(bold=True, size=13, color="0D5654")
    center = Alignment(horizontal="center")

    # --- Hoja 1: Desembolsos (detalle) --------------------------------------
    ws = book.active
    ws.title = "Desembolsos"
    cols = columnas_detalle()
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
    for reg in registros:
        fila = []
        for c in cols:
            val = reg.get(c, "")
            if c in MONEY_LABELS:
                val = _to_number(val)
            fila.append(val)
        ws.append(fila)
    # Formato de dinero y ancho de columnas.
    money_cols = [i + 1 for i, c in enumerate(cols) if c in MONEY_LABELS]
    for idx in money_cols:
        letter = get_column_letter(idx)
        for cell in ws[letter][1:]:
            cell.number_format = '#,##0.00'
    for i, c in enumerate(cols, start=1):
        letter = get_column_letter(i)
        width = min(max(len(c) + 2, 12), 40)
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    if registros:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(registros) + 1}"

    # --- Hoja 2: Por usuario -------------------------------------------------
    ws2 = book.create_sheet("Por usuario")
    por_usuario: dict[str, dict[str, Any]] = {}
    for reg in registros:
        uid = reg.get("_uid") or reg.get("Usuario") or "(sin id)"
        agg = por_usuario.setdefault(uid, {
            "Usuario": reg.get("Usuario", ""),
            "ID de usuario": uid,
            "Solicitudes": 0,
            "Monto total": 0.0,
            "Canal": reg.get("Cliente · canal", "") or reg.get("Conexión (Tipo)", ""),
            "Estado cliente": reg.get("Cliente · estado", ""),
            "Última actividad": reg.get("Cliente · última actividad", ""),
        })
        agg["Solicitudes"] += 1
        agg["Monto total"] += float(_to_number(reg.get("Monto solicitado", 0)) or 0)
    encabezado_u = ["Usuario", "ID de usuario", "Solicitudes", "Monto total",
                    "Canal", "Estado cliente", "Última actividad"]
    ws2.append(encabezado_u)
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
    for agg in sorted(por_usuario.values(), key=lambda a: a["Monto total"], reverse=True):
        ws2.append([agg[h] for h in encabezado_u])
    for cell in ws2["D"][1:]:
        cell.number_format = '#,##0.00'
    for i, c in enumerate(encabezado_u, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = min(max(len(c) + 2, 12), 32)
    ws2.freeze_panes = "A2"
    if por_usuario:
        ws2.auto_filter.ref = f"A1:{get_column_letter(len(encabezado_u))}{len(por_usuario) + 1}"

    # --- Hoja 3: Resumen -----------------------------------------------------
    ws3 = book.create_sheet("Resumen")
    total_monto = sum(float(_to_number(r.get("Monto solicitado", 0)) or 0) for r in registros)
    por_canal: dict[str, int] = {}
    por_estado: dict[str, int] = {}
    por_producto: dict[str, int] = {}
    for r in registros:
        canal = r.get("Cliente · canal") or r.get("Conexión (Tipo)") or "(desconocido)"
        por_canal[canal] = por_canal.get(canal, 0) + 1
        estado = r.get("Estado") or "(sin estado)"
        por_estado[estado] = por_estado.get(estado, 0) + 1
        prod = r.get("Producto causal") or "(sin producto)"
        por_producto[prod] = por_producto.get(prod, 0) + 1

    ws3["A1"] = "Resumen de desembolsos (retiros)"
    ws3["A1"].font = title_font
    filas_resumen: list[tuple[str, Any]] = [
        ("", ""),
        ("Generado", datetime.now().isoformat(timespec="seconds")),
        ("Origen (CSV)", str(origen)),
        ("Modo de detección", modo),
        ("Solicitudes de desembolso", len(registros)),
        ("Clientes distintos", len(por_usuario)),
        ("Monto total solicitado", round(total_monto, 2)),
        ("", ""),
        ("Por canal", ""),
    ]
    for k, v in sorted(por_canal.items(), key=lambda x: x[1], reverse=True):
        filas_resumen.append((f"   {k}", v))
    filas_resumen.append(("", ""))
    filas_resumen.append(("Por estado", ""))
    for k, v in sorted(por_estado.items(), key=lambda x: x[1], reverse=True):
        filas_resumen.append((f"   {k}", v))
    filas_resumen.append(("", ""))
    filas_resumen.append(("Por producto", ""))
    for k, v in sorted(por_producto.items(), key=lambda x: x[1], reverse=True):
        filas_resumen.append((f"   {k}", v))
    for k, v in filas_resumen:
        ws3.append([k, v])
    ws3.column_dimensions["A"].width = 34
    ws3.column_dimensions["B"].width = 40

    return book


def escribir_excel(registros: list[dict[str, Any]], salida: Path,
                   origen: Path, modo: str) -> Path:
    salida.parent.mkdir(parents=True, exist_ok=True)
    book = construir_libro(registros, origen, modo)
    book.save(salida)
    return salida


# --------------------------------------------------------------------------- #
# Descubrimiento del CSV de origen.
# --------------------------------------------------------------------------- #
def csv_mas_reciente(directorio: Path = DEFAULT_CSV_DIR) -> Path | None:
    if not directorio.is_dir():
        return None
    candidatos = sorted(
        list(directorio.glob("transacciones_producto*.csv")) or list(directorio.glob("*.csv")),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    return candidatos[0] if candidatos else None


def resolver_origen(arg: str | None) -> Path:
    if arg:
        p = Path(arg)
        if p.is_dir():
            latest = csv_mas_reciente(p)
            if latest:
                return latest
            raise SystemExit(f"No hay CSV en {p}")
        if p.exists():
            return p
        raise SystemExit(f"No existe el CSV: {p}")
    latest = csv_mas_reciente()
    if latest:
        return latest
    raise SystemExit(
        f"No hay CSV en {DEFAULT_CSV_DIR}. Ejecuta antes 'python extraccionDatos.py' "
        "o pasa la ruta del CSV como argumento."
    )


def generar(origen: Path, salida: Path, modo: str,
            productos: set[str], terminos: tuple[str, ...]) -> tuple[Path, int]:
    perfiles = cargar_usuarios()
    registros = list(iter_desembolsos(origen, perfiles, modo, productos, terminos))
    escribir_excel(registros, salida, origen, modo)
    return salida, len(registros)


def _config_extra() -> tuple[set[str], tuple[str, ...]]:
    productos = set(PAYMENT_PRODUCTS)
    extra_p = os.getenv("DESEMBOLSO_PRODUCTOS", "").strip()
    if extra_p:
        productos |= {_norm_key(x) for x in extra_p.split(",") if x.strip()}
    terminos = list(WITHDRAWAL_TERMS)
    extra_t = os.getenv("DESEMBOLSO_TERMINOS", "").strip()
    if extra_t:
        terminos += [_norm_key(x) for x in extra_t.split(",") if x.strip()]
    return productos, tuple(terminos)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Exporta a Excel las solicitudes de desembolso enriquecidas con el JSON de usuarios.",
    )
    parser.add_argument("csv", nargs="?", help="CSV de transacciones (o carpeta). Por defecto, el más reciente de descargas/.")
    parser.add_argument("--salida", "-o", default=str(DEFAULT_OUTPUT), help="Ruta del Excel de salida.")
    parser.add_argument("--modo", choices=["pagos", "amplio", "withdraw"], default="pagos",
                        help="pagos=retiros de caja (defecto); amplio=+términos de retiro; withdraw=todo 'Withdraw' (crudo).")
    parser.add_argument("--watch", action="store_true", help="Regenera en tiempo real cuando el CSV cambie.")
    parser.add_argument("--intervalo", type=int, default=20, help="Segundos entre comprobaciones en modo --watch.")
    args = parser.parse_args(argv)

    productos, terminos = _config_extra()
    salida = Path(args.salida)

    if not args.watch:
        origen = resolver_origen(args.csv)
        ruta, n = generar(origen, salida, args.modo, productos, terminos)
        print(f"OK · {n} desembolso(s) · {ruta}  (origen: {origen.name}, modo: {args.modo})")
        return 0

    # Modo tiempo real: regenera cuando cambia el mtime del CSV de origen.
    print(f"[watch] Vigilando desembolsos cada {args.intervalo}s. Ctrl+C para salir.")
    last_mtime = 0.0
    last_origen: Path | None = None
    try:
        while True:
            try:
                origen = resolver_origen(args.csv)
                mtime = origen.stat().st_mtime
                if origen != last_origen or mtime != last_mtime:
                    ruta, n = generar(origen, salida, args.modo, productos, terminos)
                    stamp = datetime.now().strftime("%H:%M:%S")
                    print(f"[{stamp}] actualizado · {n} desembolso(s) · {ruta.name}  ({origen.name})")
                    last_mtime, last_origen = mtime, origen
            except SystemExit as exc:
                print(f"[watch] {exc}")
            time.sleep(max(args.intervalo, 3))
    except KeyboardInterrupt:
        print("\n[watch] detenido.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
